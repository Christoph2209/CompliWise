"""
roster_file.py

Reads an uploaded students or staff file into plain row dicts, whether
it is a CSV export or an Excel workbook (.xlsx), and understands the
"compliance roster" layout many schools keep by hand: one row per
student, a class-placement column per core subject ("Gen Education",
"ICT", "12:1+1" ...) and a column per related service ("2X30 group of
3").

Nothing here touches the database. import_validation.py and
import_csv_data.py both read files through this module, so a file is
previewed and imported from exactly the same rows.
"""

from __future__ import annotations

import csv
import re
import zipfile
from pathlib import Path
from typing import Any, Optional, Union

from openpyxl import load_workbook

from scheduling_core import RESOURCE_ROOM, TEACHING_ASSISTANT

PathLike = Union[str, Path]

MAX_ROWS = 20_000  # far beyond any single school; guards against huge uploads

# A workbook's header row is the first row with at least this many
# filled cells, so title and subtitle lines above the table are skipped.
MIN_HEADER_CELLS = 3
HEADER_SEARCH_ROWS = 50
FIRST_NAME_HEADERS = {"first name", "first_name"}

_OLD_XLS_MAGIC = b"\xd0\xcf\x11\xe0"


class RosterFileError(Exception):
    """The file couldn't be read at all. The message is shown to the user."""


# ---------------------------------------------------------------
# Reading CSV and Excel files
# ---------------------------------------------------------------

def read_table(path: PathLike) -> tuple[list[str], list[tuple[int, dict]]]:
    """
    Read a CSV or .xlsx file into (headers, rows). Each row comes with
    its spreadsheet row number, so problems can be reported against the
    line the school sees in Excel. Header names are stripped and every
    cell is text.

    The file type is decided by its contents, not its name. Raises
    RosterFileError when the file can't be read or is too large.
    """
    path = Path(path)
    with path.open("rb") as f:
        magic = f.read(4)
    if magic == _OLD_XLS_MAGIC:
        raise RosterFileError(
            "Old-style .xls workbooks aren't supported. Save the file as .xlsx or CSV and upload it again."
        )
    if zipfile.is_zipfile(path):
        return _read_workbook(path)
    return _read_csv(path)


def _read_csv(path: Path) -> tuple[list[str], list[tuple[int, dict]]]:
    try:
        # utf-8-sig so an Excel BOM doesn't hide the first column name.
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            headers = [h.strip() for h in (reader.fieldnames or [])]
            rows = []
            for row_num, row in enumerate(reader, start=2):
                rows.append((row_num, {(k or "").strip(): v for k, v in row.items()}))
                if len(rows) > MAX_ROWS:
                    raise RosterFileError(f"File has more than {MAX_ROWS:,} rows.")
            return headers, rows
    except UnicodeDecodeError:
        raise RosterFileError(
            "File isn't UTF-8 text. In Excel, use Save As > 'CSV UTF-8 (Comma delimited)', "
            "or upload the .xlsx workbook itself."
        ) from None
    except csv.Error as error:
        raise RosterFileError(f"Couldn't read the file as CSV: {error}") from None


def _cell_text(value: Any) -> str:
    """A workbook cell as the text a CSV would hold (6.0 -> "6")."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _sheet_table(sheet) -> tuple[list[str], list[tuple[int, dict]]]:
    """(headers, rows) of one worksheet; ([], []) when it has no header row."""
    columns: dict[int, str] = {}
    rows: list[tuple[int, dict]] = []

    for row_num, cells in enumerate(sheet.iter_rows(min_row=1, values_only=True), start=1):
        texts = [_cell_text(cell) for cell in cells]
        if not columns:
            filled = {i: t.strip() for i, t in enumerate(texts) if t.strip()}
            if len(filled) >= MIN_HEADER_CELLS:
                columns = filled
            elif row_num >= HEADER_SEARCH_ROWS:
                break
            continue
        if not any(t.strip() for t in texts):
            continue
        rows.append((row_num, {
            name: texts[i] if i < len(texts) else "" for i, name in columns.items()
        }))
        if len(rows) > MAX_ROWS:
            raise RosterFileError(f"File has more than {MAX_ROWS:,} rows.")

    return list(columns.values()), rows


def _read_workbook(path: Path) -> tuple[list[str], list[tuple[int, dict]]]:
    # Opened as a stream because openpyxl refuses a path that doesn't end
    # in .xlsx, and uploads are saved under a fixed name.
    with path.open("rb") as stream:
        try:
            workbook = load_workbook(stream, read_only=True, data_only=True)
        except Exception:
            # openpyxl raises several unrelated error types for a damaged
            # workbook or a zip file that isn't a workbook at all.
            raise RosterFileError(
                "Couldn't read the file as an Excel workbook. Save it as .xlsx or CSV and upload it again."
            ) from None

        try:
            # Workbooks often carry extra sheets (e.g. the bell schedule);
            # the roster is the first one with a name column.
            fallback = None
            for sheet in workbook.worksheets:
                headers, rows = _sheet_table(sheet)
                if not headers:
                    continue
                if FIRST_NAME_HEADERS & {h.lower() for h in headers}:
                    return headers, rows
                if fallback is None:
                    fallback = (headers, rows)
            return fallback or ([], [])
        finally:
            workbook.close()


# ---------------------------------------------------------------
# Compliance-roster columns
# ---------------------------------------------------------------

# Column headers are matched case-insensitively.
PLACEMENT_COLUMNS = {
    "ela": "ELA",
    "math": "Math",
    "science": "Science",
    "social studies": "Social Studies",
}
SESSION_SERVICE_COLUMNS = {
    "speech": "Speech",
    "ot": "OT",
    "pt": "PT",
    "counseling": "Counseling",
    "resource room": RESOURCE_ROOM,
    "setss": RESOURCE_ROOM,
}
TEACHING_ASSISTANT_COLUMN = "teaching assistant"
ENL_LEVEL_COLUMN = "enl"

GENERAL_ED = "General Ed"
ICT = "ICT"
_GENERAL_ED_SPELLINGS = {"geneducation", "generaleducation", "gened", "generaled", "general", "ge"}
# "12:1+1", "12:1", "8:1+1" ... students : teachers (+ paraprofessionals).
_CLASS_RATIO_RE = re.compile(r"(\d+)\s*:\s*(\d+)(?:\s*[+:]\s*(\d+))?")

# A class placement is every period of that subject. The roster doesn't
# say how long a period is, so one period a day of this length is assumed.
CLASS_PERIOD_MINUTES = 45
SCHOOL_DAYS_PER_WEEK = 5

# "2X30 group of 3" = 2 sessions a week of 30 minutes; "5Xweek" has no length.
_SESSIONS_RE = re.compile(r"(\d+)\s*x\s*(\d+)?", re.IGNORECASE)
_GROUP_RE = re.compile(r"group\s*(?:of)?\s*(\d+)", re.IGNORECASE)
_INDIVIDUAL_RE = re.compile(r"one\s*to\s*one|1\s*:\s*1|individual", re.IGNORECASE)

# Weekly ENL minutes NYS CR Part 154 requires at each proficiency level
# in grades K-8. The roster only gives the level.
ENL_WEEKLY_MINUTES = {
    "entering": 360,
    "emerging": 360,
    "transitioning": 180,
    "expanding": 180,
    "commanding": 90,
}


def _by_lowercase_header(row: dict) -> dict:
    return {k.strip().lower(): str(v or "").strip() for k, v in row.items() if isinstance(k, str)}


def parse_placement(text: str) -> Optional[str]:
    """
    A class-placement cell as a canonical name, however it was typed:

    "gen Education"      -> "General Ed"
    "ICt "               -> "ICT"
    "12:1+1"             -> "12:1+1"
    "Special class12:1"  -> "12:1"

    None when the text isn't a placement this importer knows.
    """
    squeezed = re.sub(r"[\s.\-]", "", text).lower()
    if squeezed in _GENERAL_ED_SPELLINGS:
        return GENERAL_ED
    if squeezed == "ict":
        return ICT
    ratio = _CLASS_RATIO_RE.search(text)
    if ratio:
        students, teachers, paras = ratio.groups()
        return f"{students}:{teachers}+{paras}" if paras else f"{students}:{teachers}"
    return None


def parse_sessions(text: str) -> Optional[dict]:
    """
    A related-service cell as sessions, minutes and group size:

    "2X30 group of 3"  -> {"sessions_per_week": 2, "minutes_per_session": 30, "group_size": 3}
    "1X30 one to one"  -> {"sessions_per_week": 1, "minutes_per_session": 30, "group_size": 1}
    "5Xweek"           -> {"sessions_per_week": 5, "minutes_per_session": None, "group_size": None}

    None when no session count can be read.
    """
    match = _SESSIONS_RE.match(text.strip())
    if not match or int(match.group(1)) < 1:
        return None

    group = _GROUP_RE.search(text)
    if group:
        group_size = int(group.group(1))
    elif _INDIVIDUAL_RE.search(text[match.end():]):
        group_size = 1
    else:
        group_size = None

    return {
        "sessions_per_week": int(match.group(1)),
        "minutes_per_session": int(match.group(2)) if match.group(2) else None,
        "group_size": group_size,
    }


def roster_services(row: dict) -> tuple[list[dict], list[tuple[str, str]]]:
    """
    The services one student row lists in compliance-roster columns, as
    (services, warnings). Each service is a dict shaped like an
    `iep_services` JSON entry, plus the details a roster adds
    (minutes_per_session, subject_area, is_pullout, note). Each warning
    is (column, message) for a cell that will import with an assumption.

    A row from a file without these columns gives ([], []).
    """
    cells = _by_lowercase_header(row)
    services: list[dict] = []
    warnings: list[tuple[str, str]] = []

    for column, subject in PLACEMENT_COLUMNS.items():
        if column not in cells:
            continue
        text = cells[column]
        if not text:
            warnings.append((subject, f"No {subject} placement is listed; none will be recorded."))
            continue
        placement = parse_placement(text)
        if placement is None:
            warnings.append((subject, f"{subject} placement '{text}' isn't recognized; none will be recorded."))
        elif placement != GENERAL_ED:
            services.append({
                "service_type": placement,
                "subject_area": subject,
                "sessions_per_week": SCHOOL_DAYS_PER_WEEK,
                "minutes_per_session": CLASS_PERIOD_MINUTES,
                "is_pullout": False,
                "note": f"Class placement for {subject}. Minutes assume one "
                        f"{CLASS_PERIOD_MINUTES}-minute period a day.",
            })

    for column, service_type in SESSION_SERVICE_COLUMNS.items():
        text = cells.get(column)
        if not text:
            continue
        sessions = parse_sessions(text)
        if sessions is None:
            warnings.append((service_type, f"{service_type}: '{text}' not understood (expected like "
                                           "'2X30 group of 3'); 1 session/week will be assumed."))
            sessions = {}
        services.append({"service_type": service_type, **sessions})

    assistant = cells.get(TEACHING_ASSISTANT_COLUMN)
    if assistant:
        services.append({
            "service_type": TEACHING_ASSISTANT,
            "sessions_per_week": SCHOOL_DAYS_PER_WEEK,
            "minutes_per_session": CLASS_PERIOD_MINUTES,
            "is_pullout": False,
            "note": f"{assistant}. Stays with the student, so the minutes are a placeholder.",
        })

    return services, warnings


def roster_enl_level(row: dict) -> Optional[str]:
    """The roster's ENL proficiency level, lowercased ("transitioning"); None if blank or absent."""
    return _by_lowercase_header(row).get(ENL_LEVEL_COLUMN, "").lower() or None
